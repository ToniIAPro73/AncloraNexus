import os
import tempfile
import pytest
from src.models.conversion import conversion_engine
from fpdf import FPDF
from PIL import Image
from docx import Document


def create_sample_file(ext: str, path: str):
    if ext == 'txt':
        with open(path, 'w', encoding='utf-8') as f:
            f.write('hola mundo')
    elif ext == 'pdf':
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font('Arial', size=12)
        pdf.cell(0, 10, 'hola pdf', ln=1)
        pdf.output(path)
    elif ext in ('jpg', 'png', 'gif'):
        img = Image.new('RGB', (50, 50), color='red')
        img.save(path)
    elif ext == 'svg':
        svg_content = (
            '<svg xmlns="http://www.w3.org/2000/svg" width="50" height="50">'
            '<rect width="50" height="50" fill="red"/></svg>'
        )
        with open(path, 'w', encoding='utf-8') as f:
            f.write(svg_content)
    elif ext == 'docx':
        doc = Document()
        doc.add_paragraph('hola docx')
        doc.save(path)
    elif ext == 'doc':
        with open(path, 'w', encoding='utf-8') as f:
            f.write('hola doc')
    else:
        raise ValueError('unsupported ext')


def test_plugins_discovered():
    """Ensure plugins are loaded into the registry."""
    assert ('txt', 'pdf') in conversion_engine.conversion_methods


@pytest.mark.parametrize('source_ext,target_ext', [
    ('txt', 'doc'), ('txt', 'docx'), ('txt', 'pdf'), ('txt', 'odt'), ('txt', 'tex'),
    ('pdf', 'jpg'), ('pdf', 'png'), ('pdf', 'gif'), ('pdf', 'txt'),
    ('jpg', 'png'), ('jpg', 'pdf'), ('jpg', 'gif'),
    ('png', 'jpg'), ('png', 'pdf'), ('png', 'gif'), ('png', 'webp'),
    ('gif', 'jpg'), ('gif', 'png'), ('gif', 'pdf'), ('gif', 'mp4'),
    ('doc', 'pdf'), ('doc', 'txt'), ('doc', 'html'),
    ('docx', 'pdf'), ('docx', 'txt'), ('docx', 'html'),
    ('svg', 'png')
])
def test_conversion_engine(source_ext, target_ext):
    with tempfile.TemporaryDirectory() as tmp:
        input_path = os.path.join(tmp, f'in.{source_ext}')
        output_path = os.path.join(tmp, f'out.{target_ext}')
        create_sample_file(source_ext, input_path)
        result, msg = conversion_engine.convert_file(input_path, output_path, source_ext, target_ext)
        assert result, msg
        assert os.path.exists(output_path)
